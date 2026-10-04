"""Tests for Digital PDF and DOCX Text Extractor (TASK-11.2).

Verifies:
- PDF digital text extraction with page tracking, boundaries, and encryption handling.
- DOCX text extraction with paragraph sequence, headings, and table formatting.
- Text normalization preserving scientific notation, mathematical symbols, and layout.
- Scanned / textless document detection without silent failure.
- Database persistence and ownership isolation over API routes.
"""

from datetime import datetime, timezone
import io
from pathlib import Path
from typing import Optional
from unittest.mock import AsyncMock, MagicMock
import uuid

import docx
import httpx
import pypdf
import pytest
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.security import create_access_token
from app.main import create_app
from app.modules.auth.models import User
from app.modules.documents.extractor import (
    DocumentExtractorService,
    DocumentFileNotFoundError,
    ExtractionStatusEnum,
    TextNormalizer,
    get_extractor_service,
)
from app.modules.documents.models import UploadedDocument
from app.modules.documents.schemas import DocumentStatusEnum
from app.modules.documents.storage import DocumentStorageService, get_storage_service
from app.modules.experiments.models import Experiment, PreparationChecklist


# ---------------------------------------------------------------------------
# Synthetic Test Fixtures Generators
# ---------------------------------------------------------------------------


def make_synthetic_pdf(
    page_texts: list[str],
    encrypted: bool = False,
    password: str = "secret123",
) -> bytes:
    """Create a valid in-memory PDF with the provided page texts."""
    writer = pypdf.PdfWriter()
    for text in page_texts:
        if not text:
            # Blank page
            writer.add_blank_page(width=612, height=792)
            continue

        clean_text = text.replace("(", "").replace(")", "")
        stream_str = f"BT /F1 12 Tf 72 712 Td ({clean_text}) Tj ET"
        stream_bytes = stream_str.encode("latin-1", errors="replace")
        length = len(stream_bytes)
        raw_pdf = (
            b"%PDF-1.4\n"
            b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
            b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
            b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
            + f"4 0 obj << /Length {length} >> stream\n".encode("latin-1")
            + stream_bytes
            + b"\nendstream\nendobj\n"
            b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
            b"xref\n0 6\n0000000000 65535 f \n0000000009 00000 n \n0000000058 00000 n \n"
            b"0000000115 00000 n \n0000000266 00000 n \n0000000360 00000 n \n"
            b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n441\n%%EOF\n"
        )
        r = pypdf.PdfReader(io.BytesIO(raw_pdf))
        writer.add_page(r.pages[0])

    if encrypted:
        writer.encrypt(password)

    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_synthetic_docx(
    title: str = "Experiment 1",
    sections: Optional[list[tuple[str, str]]] = None,
    table_data: Optional[list[list[str]]] = None,
) -> bytes:
    """Create a valid in-memory DOCX document with headings, paragraphs, and tables."""
    doc = docx.Document()
    if title:
        doc.add_heading(title, level=1)

    if sections:
        for heading, body in sections:
            doc.add_heading(heading, level=2)
            doc.add_paragraph(body)

    if table_data and len(table_data) > 0:
        tbl = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                tbl.cell(r_idx, c_idx).text = val

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# Test Fixtures & Mock App Builder
# ---------------------------------------------------------------------------


@pytest.fixture
def sample_user_id() -> uuid.UUID:
    return uuid.uuid4()


@pytest.fixture
def sample_user(sample_user_id: uuid.UUID) -> User:
    now = datetime.now(timezone.utc)
    return User(
        id=sample_user_id,
        email="labuser@university.edu",
        password_hash="$argon2id$v=19$m=65536,t=3,p=4$somehash",
        full_name="Lab Student",
        university="University",
        is_active=True,
        created_at=now,
        updated_at=now,
    )


@pytest.fixture
def auth_token(sample_user_id: uuid.UUID) -> str:
    return create_access_token(sample_user_id)


@pytest.fixture
def auth_headers(auth_token: str) -> dict[str, str]:
    return {"Authorization": f"Bearer {auth_token}"}


def create_sample_experiment(
    user_id: uuid.UUID,
    exp_id: Optional[uuid.UUID] = None,
    title: str = "Refraction of Light",
    has_manual_file: bool = False,
    file_name: Optional[str] = None,
) -> Experiment:
    """Helper to instantiate an Experiment ORM model with checklist."""
    now = datetime.now(timezone.utc)
    actual_exp_id = exp_id or uuid.uuid4()
    exp = Experiment(
        id=actual_exp_id,
        user_id=user_id,
        title=title,
        subject="Physics",
        experiment_number="PHY-102",
        course_semester="Semester 2",
        creation_method="manual",
        has_manual_file=has_manual_file,
        file_name=file_name,
        status="ready",
        description="Measurement of refractive index of prism.",
        objective="Determine refractive index of prism material.",
        theory="Snell's Law and angle of minimum deviation.",
        apparatus="Spectrometer, prism, sodium lamp.",
        procedure="1. Mount prism. 2. Measure angle of prism. 3. Measure minimum deviation.",
        observations="Table of deviation angles.",
        calculations="μ = sin((A+D)/2) / sin(A/2).",
        precautions="Do not touch optical surfaces.",
        created_at=now,
        updated_at=now,
    )
    chk = PreparationChecklist(
        id=uuid.uuid4(),
        experiment_id=actual_exp_id,
        items={"objective": True, "theory": True, "apparatus": True, "procedure": True, "precautions": True},
        created_at=now,
        updated_at=now,
        experiment=exp,
    )
    exp.checklist = chk
    exp.documents = []
    return exp


def build_test_client(
    sample_user: User,
    experiment: Optional[Experiment] = None,
    document: Optional[UploadedDocument] = None,
    storage_service: Optional[DocumentStorageService] = None,
    extractor_service: Optional[DocumentExtractorService] = None,
) -> tuple[httpx.AsyncClient, AsyncMock]:
    """Build test client with mocked DB session and dependency overrides."""
    db_session = AsyncMock(spec=AsyncSession)
    db_session.add = MagicMock()
    db_session.delete = AsyncMock()
    db_session.commit = AsyncMock()
    db_session.rollback = AsyncMock()
    db_session.refresh = AsyncMock()

    async def fake_execute(statement, *args, **kwargs):
        stmt_str = str(statement).lower()
        res = MagicMock()

        if "from users" in stmt_str:
            res.scalar_one_or_none.return_value = sample_user
        elif "from experiments" in stmt_str:
            res.scalar_one_or_none.return_value = experiment
        elif "from uploaded_documents" in stmt_str:
            res.scalar_one_or_none.return_value = document
        else:
            res.scalar_one_or_none.return_value = None
            res.scalars.return_value.all.return_value = []

        return res

    db_session.execute = AsyncMock(side_effect=fake_execute)

    app = create_app()
    app.dependency_overrides[get_db] = lambda: db_session
    if storage_service is not None:
        app.dependency_overrides[get_storage_service] = lambda: storage_service
    if extractor_service is not None:
        app.dependency_overrides[get_extractor_service] = lambda: extractor_service

    transport = httpx.ASGITransport(app=app)
    client = httpx.AsyncClient(transport=transport, base_url="http://test")
    return client, db_session


# ---------------------------------------------------------------------------
# Text Normalization Tests
# ---------------------------------------------------------------------------


def test_normalizer_line_endings_and_crlf():
    """Verify that CRLF and CR are normalized to standard LF."""
    raw = "Line 1\r\nLine 2\rLine 3\nLine 4"
    normalized = TextNormalizer.normalize(raw)
    assert "\r" not in normalized
    assert normalized == "Line 1\nLine 2\nLine 3\nLine 4"


def test_normalizer_removes_excess_whitespace_and_blank_lines():
    """Verify that multiple spaces within lines and 3+ blank lines are cleanly collapsed."""
    raw = "Section 1    has    wide    spaces\n\n\n\n\nSection 2 follows with trailing spaces   \n\n\nSection 3"
    normalized = TextNormalizer.normalize(raw)
    assert "Section 1 has wide spaces" in normalized
    assert "trailing spaces" in normalized
    assert "\n\n\n" not in normalized
    assert "Section 1 has wide spaces\n\nSection 2 follows with trailing spaces\n\nSection 3" == normalized


def test_normalizer_preserves_scientific_notation_and_units():
    """Verify mathematical symbols, superscripts, Greek letters, and formulas remain uncorrupted."""
    scientific_text = (
        "Calculations:\n"
        "g = (4 * π² * L) / T²\n"
        "Value: 9.81 ± 0.05 m/s²\n"
        "Viscosity: η = 1.002 x 10^-3 Pa·s at 20°C\n"
        "Resistance: R = 100 Ω, ΔV = 5.0 V, Current I = 50 mA"
    )
    normalized = TextNormalizer.normalize(scientific_text)
    assert "g = (4 * π² * L) / T²" in normalized
    assert "9.81 ± 0.05 m/s²" in normalized
    assert "η = 1.002 x 10^-3 Pa·s at 20°C" in normalized
    assert "R = 100 Ω, ΔV = 5.0 V, Current I = 50 mA" in normalized


def test_normalizer_preserves_markdown_tables():
    """Verify that markdown tables retain formatting and row integrity."""
    table_text = (
        "| Parameter | Value | Unit |\n"
        "| --- | --- | --- |\n"
        "| Length | 100.5 | cm |\n"
        "| Mass | 50.2 | g |"
    )
    normalized = TextNormalizer.normalize(table_text)
    assert "| Parameter | Value | Unit |" in normalized
    assert "| Length | 100.5 | cm |" in normalized


# ---------------------------------------------------------------------------
# PDF Extraction Tests
# ---------------------------------------------------------------------------


def test_extract_valid_digital_pdf():
    """Verify clean text extraction from a valid digital PDF."""
    pdf_bytes = make_synthetic_pdf([
        "Physics Lab Manual: Measurement of Gravitational Acceleration using Simple Pendulum with length L = 100cm."
    ])
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(pdf_bytes, "manual.pdf")

    assert result.status == ExtractionStatusEnum.SUCCESS
    assert result.source_file_type == "pdf"
    assert result.page_count == 1
    assert "Simple Pendulum" in result.extracted_text
    assert "L = 100cm" in result.extracted_text
    assert len(result.pages) == 1
    assert result.pages[0].has_text is True


def test_extract_pdf_multi_page_preserves_boundaries():
    """Verify extraction across multiple pages preserves page metadata and ordering."""
    p1 = "Page 1: Theory and Objective of Ohm Law Verification in DC Circuits."
    p2 = "Page 2: Apparatus Required: Voltmeter, Ammeter, Rheostat, DC Power Supply."
    pdf_bytes = make_synthetic_pdf([p1, p2])

    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(pdf_bytes, "manual.pdf")

    assert result.status == ExtractionStatusEnum.SUCCESS
    assert result.page_count == 2
    assert len(result.pages) == 2
    assert result.pages[0].page_number == 1
    assert "Page 1: Theory" in result.pages[0].text
    assert result.pages[1].page_number == 2
    assert "Page 2: Apparatus" in result.pages[1].text
    assert "Page 1: Theory" in result.extracted_text
    assert "Page 2: Apparatus" in result.extracted_text


def test_extract_pdf_with_empty_page():
    """Verify handling when a page has no text."""
    p1 = "Page 1: Substantial text content for experiment procedure."
    p2 = ""  # Blank page
    pdf_bytes = make_synthetic_pdf([p1, p2])

    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(pdf_bytes, "manual.pdf")

    assert result.status in (ExtractionStatusEnum.SUCCESS, ExtractionStatusEnum.SUCCESS_WITH_WARNINGS)
    assert result.page_count == 2
    assert result.pages[0].has_text is True
    assert result.pages[1].has_text is False
    assert result.pages[1].text == ""


def test_extract_pdf_textless_or_scanned_flags_no_text():
    """Verify that a completely empty/scanned PDF is flagged with NO_TEXT_FOUND without crashing."""
    pdf_bytes = make_synthetic_pdf(["", ""])
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(pdf_bytes, "scanned.pdf")

    assert result.status == ExtractionStatusEnum.NO_TEXT_FOUND
    assert result.extracted_text == ""
    assert result.character_count == 0
    assert any("scanned" in w.lower() or "no extractable digital text" in w.lower() for w in result.warnings)


def test_extract_pdf_encrypted_fails_safely():
    """Verify that encrypted password-protected PDFs are detected with ENCRYPTED status."""
    pdf_bytes = make_synthetic_pdf(["Secret manual content."], encrypted=True, password="mypassword")
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(pdf_bytes, "protected.pdf")

    assert result.status == ExtractionStatusEnum.ENCRYPTED
    assert result.extracted_text == ""
    assert "password" in result.failure_reason.lower() or "encrypted" in result.failure_reason.lower()


def test_extract_pdf_corrupted_fails_safely():
    """Verify corrupted PDF bytes produce CORRUPTED status rather than crashing."""
    corrupted_bytes = b"%PDF-1.4\ncorrupted random bytes that do not form a valid pdf xref\n%%EOF"
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(corrupted_bytes, "bad.pdf")

    assert result.status == ExtractionStatusEnum.CORRUPTED
    assert result.failure_reason is not None


# ---------------------------------------------------------------------------
# DOCX Extraction Tests
# ---------------------------------------------------------------------------


def test_extract_valid_docx():
    """Verify clean text and table extraction from a valid DOCX."""
    docx_bytes = make_synthetic_docx(
        title="Chemistry Lab: Acid-Base Titration",
        sections=[
            ("Objective", "Determine the molarity of standard HCl solution."),
            ("Apparatus", "Burette, pipette, conical flask, phenolphthalein indicator."),
        ],
        table_data=[
            ["Trial", "Initial Burette (mL)", "Final Burette (mL)", "Volume Used (mL)"],
            ["1", "0.0", "20.1", "20.1"],
            ["2", "0.0", "20.0", "20.0"],
        ],
    )
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(docx_bytes, "manual.docx")

    assert result.status == ExtractionStatusEnum.SUCCESS
    assert result.source_file_type == "docx"
    assert "# Chemistry Lab: Acid-Base Titration" in result.extracted_text
    assert "## Objective" in result.extracted_text
    assert "## Apparatus" in result.extracted_text
    assert "| Trial | Initial Burette (mL) | Final Burette (mL) | Volume Used (mL) |" in result.extracted_text
    assert "| 1 | 0.0 | 20.1 | 20.1 |" in result.extracted_text


def test_extract_docx_empty_document():
    """Verify empty DOCX document returns NO_TEXT_FOUND status."""
    docx_bytes = make_synthetic_docx(title="")
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(docx_bytes, "empty.docx")

    assert result.status == ExtractionStatusEnum.NO_TEXT_FOUND
    assert result.character_count == 0


def test_extract_docx_corrupted_fails_safely():
    """Verify corrupted DOCX content returns CORRUPTED status safely."""
    corrupted_bytes = b"PK\x03\x04corrupted zip payload that is not a real openxml document"
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(corrupted_bytes, "bad.docx")

    assert result.status == ExtractionStatusEnum.CORRUPTED
    assert result.failure_reason is not None


def test_extract_unsupported_format():
    """Verify unsupported formats return UNSUPPORTED status."""
    txt_bytes = b"Plain text file content."
    extractor = DocumentExtractorService()
    result = extractor.extract_from_bytes(txt_bytes, "manual.txt")

    assert result.status == ExtractionStatusEnum.UNSUPPORTED
    assert "unsupported" in result.failure_reason.lower()


def test_extract_missing_file_from_path():
    """Verify that a nonexistent file path returns FAILED status without crashing."""
    extractor = DocumentExtractorService()
    result = extractor.extract_from_path("/tmp/nonexistent_manual_file_12345.pdf")

    assert result.status == ExtractionStatusEnum.FAILED
    assert "does not exist" in result.failure_reason


# ---------------------------------------------------------------------------
# Database Persistence & API Route Integration Tests
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_extract_service_persists_to_uploaded_document(
    sample_user: User,
    tmp_path: Path,
):
    """Verify that extract_document_and_persist updates UploadedDocument in the database."""
    storage = DocumentStorageService(base_dir=tmp_path)
    extractor = DocumentExtractorService(storage_service=storage)

    # 1. Create a physical test PDF file in storage
    pdf_bytes = make_synthetic_pdf([
        "Experiment Title: Boyle Law Verification at Room Temperature T = 298K."
    ])
    exp_id = uuid.uuid4()
    user_dir = tmp_path / str(sample_user.id) / str(exp_id)
    user_dir.mkdir(parents=True, exist_ok=True)
    target_file = user_dir / "stored_manual.pdf"
    target_file.write_bytes(pdf_bytes)

    relative_path = f"{sample_user.id}/{exp_id}/stored_manual.pdf"

    # 2. Instantiate UploadedDocument model
    doc_id = uuid.uuid4()
    now = datetime.now(timezone.utc)
    doc = UploadedDocument(
        id=doc_id,
        user_id=sample_user.id,
        experiment_id=exp_id,
        file_name="boyle_manual.pdf",
        file_size_bytes=len(pdf_bytes),
        mime_type="application/pdf",
        storage_path=relative_path,
        status=DocumentStatusEnum.PENDING.value,
        extracted_text=None,
        extracted_data={},
        created_at=now,
        updated_at=now,
    )

    # 3. Create mock AsyncSession
    mock_db = AsyncMock(spec=AsyncSession)
    mock_res = MagicMock()
    mock_res.scalar_one_or_none.return_value = doc
    mock_db.execute = AsyncMock(return_value=mock_res)
    mock_db.commit = AsyncMock()
    mock_db.refresh = AsyncMock()

    # 4. Execute extraction and persistence
    result = await extractor.extract_document_and_persist(
        db=mock_db,
        document_id=doc_id,
        user_id=sample_user.id,
    )

    assert result.status == ExtractionStatusEnum.SUCCESS
    assert "Boyle Law Verification" in result.extracted_text

    # 5. Check model fields were updated
    assert doc.status == DocumentStatusEnum.COMPLETED.value
    assert "Boyle Law Verification" in doc.extracted_text
    assert doc.extracted_data.get("page_count") == 1
    assert doc.extracted_data.get("extraction_status") == "success"
    assert doc.error_message is None
    assert mock_db.commit.called


@pytest.mark.asyncio
async def test_extract_endpoint_success_pdf(
    sample_user: User,
    auth_headers: dict[str, str],
    tmp_path: Path,
):
    """Verify POST /api/v1/experiments/{id}/extract-text on an uploaded PDF manual."""
    exp = create_sample_experiment(sample_user.id, has_manual_file=True, file_name="prism_manual.pdf")
    storage = DocumentStorageService(base_dir=tmp_path)
    extractor = DocumentExtractorService(storage_service=storage)

    pdf_bytes = make_synthetic_pdf([
        "Title: Refraction through Glass Prism. Angle of deviation D = (n - 1) * A."
    ])
    user_dir = tmp_path / str(sample_user.id) / str(exp.id)
    user_dir.mkdir(parents=True, exist_ok=True)
    (user_dir / "prism.pdf").write_bytes(pdf_bytes)

    rel_path = f"{sample_user.id}/{exp.id}/prism.pdf"
    doc = UploadedDocument(
        id=uuid.uuid4(),
        user_id=sample_user.id,
        experiment_id=exp.id,
        file_name="prism_manual.pdf",
        file_size_bytes=len(pdf_bytes),
        mime_type="application/pdf",
        storage_path=rel_path,
        status=DocumentStatusEnum.PENDING.value,
        extracted_text=None,
        extracted_data={},
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )
    exp.documents = [doc]

    client, _ = build_test_client(
        sample_user=sample_user,
        experiment=exp,
        document=doc,
        storage_service=storage,
        extractor_service=extractor,
    )

    resp = await client.post(
        f"/api/v1/experiments/{exp.id}/extract-text",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["sourceFileType"] == "pdf"
    assert data["status"] == "success"
    assert "Refraction through Glass Prism" in data["extractedText"]
    assert data["pageCount"] == 1


@pytest.mark.asyncio
async def test_extract_endpoint_success_docx(
    sample_user: User,
    auth_headers: dict[str, str],
    tmp_path: Path,
):
    """Verify POST /api/v1/experiments/{id}/extract-text on an uploaded DOCX manual."""
    exp = create_sample_experiment(sample_user.id, has_manual_file=True, file_name="photosynthesis.docx")
    storage = DocumentStorageService(base_dir=tmp_path)
    extractor = DocumentExtractorService(storage_service=storage)

    docx_bytes = make_synthetic_docx(
        title="Biology Experiment: Photosynthesis Rate",
        sections=[("Objective", "Measure oxygen bubble production under light.")],
    )
    user_dir = tmp_path / str(sample_user.id) / str(exp.id)
    user_dir.mkdir(parents=True, exist_ok=True)
    (user_dir / "photo.docx").write_bytes(docx_bytes)

    rel_path = f"{sample_user.id}/{exp.id}/photo.docx"
    doc = UploadedDocument(
        id=uuid.uuid4(),
        user_id=sample_user.id,
        experiment_id=exp.id,
        file_name="photosynthesis.docx",
        file_size_bytes=len(docx_bytes),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        storage_path=rel_path,
        status=DocumentStatusEnum.PENDING.value,
        extracted_text=None,
        extracted_data={},
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )
    exp.documents = [doc]

    client, _ = build_test_client(
        sample_user=sample_user,
        experiment=exp,
        document=doc,
        storage_service=storage,
        extractor_service=extractor,
    )

    resp = await client.post(
        f"/api/v1/experiments/{exp.id}/extract-text",
        headers=auth_headers,
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["sourceFileType"] == "docx"
    assert data["status"] == "success"
    assert "# Biology Experiment: Photosynthesis Rate" in data["extractedText"]
    assert "## Objective" in data["extractedText"]


@pytest.mark.asyncio
async def test_extract_endpoint_unauthenticated_rejected(
    sample_user: User,
    tmp_path: Path,
):
    """Verify that unauthenticated extraction requests are rejected with 401."""
    exp = create_sample_experiment(sample_user.id)
    client, _ = build_test_client(sample_user=sample_user, experiment=exp)

    resp = await client.post(
        f"/api/v1/experiments/{exp.id}/extract-text"
    )
    assert resp.status_code == 401


@pytest.mark.asyncio
async def test_extract_endpoint_other_user_experiment_returns_404(
    sample_user: User,
    auth_headers: dict[str, str],
):
    """Verify that requesting extraction on another user's experiment returns 404."""
    other_user_id = uuid.uuid4()
    other_exp = create_sample_experiment(other_user_id)

    # When queried with current_user.id, the experiment query returns None
    client, _ = build_test_client(sample_user=sample_user, experiment=None)

    resp = await client.post(
        f"/api/v1/experiments/{other_exp.id}/extract-text",
        headers=auth_headers,
    )
    assert resp.status_code == 404
    assert "not found" in resp.json()["detail"].lower()


@pytest.mark.asyncio
async def test_extract_endpoint_no_document_uploaded_returns_404(
    sample_user: User,
    auth_headers: dict[str, str],
):
    """Verify that calling extract-text on an experiment without a manual returns 404."""
    exp = create_sample_experiment(sample_user.id, has_manual_file=False)
    exp.documents = []  # No documents attached

    client, _ = build_test_client(sample_user=sample_user, experiment=exp)

    resp = await client.post(
        f"/api/v1/experiments/{exp.id}/extract-text",
        headers=auth_headers,
    )
    assert resp.status_code == 404
    assert "No uploaded laboratory manual found" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_extract_endpoint_missing_storage_file_handles_gracefully(
    sample_user: User,
    auth_headers: dict[str, str],
    tmp_path: Path,
):
    """Verify that a document whose file was deleted on disk returns 404 cleanly."""
    exp = create_sample_experiment(sample_user.id, has_manual_file=True, file_name="missing.pdf")
    storage = DocumentStorageService(base_dir=tmp_path)
    extractor = DocumentExtractorService(storage_service=storage)

    # Record document pointing to nonexistent file
    doc = UploadedDocument(
        id=uuid.uuid4(),
        user_id=sample_user.id,
        experiment_id=exp.id,
        file_name="missing.pdf",
        file_size_bytes=1024,
        mime_type="application/pdf",
        storage_path=f"{sample_user.id}/{exp.id}/missing.pdf",
        status=DocumentStatusEnum.PENDING.value,
        extracted_text=None,
        extracted_data={},
        created_at=datetime.now(timezone.utc),
        updated_at=datetime.now(timezone.utc),
    )
    exp.documents = [doc]

    client, _ = build_test_client(
        sample_user=sample_user,
        experiment=exp,
        document=doc,
        storage_service=storage,
        extractor_service=extractor,
    )

    resp = await client.post(
        f"/api/v1/experiments/{exp.id}/extract-text",
        headers=auth_headers,
    )
    assert resp.status_code == 404
    assert "not found in storage" in resp.json()["detail"]
